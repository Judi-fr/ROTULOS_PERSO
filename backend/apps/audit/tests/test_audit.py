"""Tests de apps.audit: modelo inmutable, servicio de registro y endpoints
de lectura. Los tests de integración (qué vistas de otras apps disparan qué
acción) viven en los tests de esas apps (accounts, orders)."""

import csv
import io
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APITestCase

from apps.accounts.models import GroupRolePermission, RolePermission

from ..models import AuditLog
from ..services import record

User = get_user_model()
STRONG_PASSWORD = "Rotulos2026"


class AuditLogModelTests(APITestCase):
    def test_es_inmutable_no_se_puede_modificar(self):
        log = AuditLog.objects.create(
            category=AuditLog.Category.AUTH,
            action=AuditLog.Action.AUTH_LOGIN_SUCCESS,
            actor_email="a@example.com",
        )
        log.actor_email = "otro@example.com"
        with self.assertRaises(ValueError):
            log.save()

    def test_es_inmutable_no_se_puede_borrar(self):
        log = AuditLog.objects.create(
            category=AuditLog.Category.AUTH,
            action=AuditLog.Action.AUTH_LOGIN_SUCCESS,
            actor_email="a@example.com",
        )
        with self.assertRaises(ValueError):
            log.delete()
        self.assertEqual(AuditLog.objects.count(), 1)


class RecordServiceTests(APITestCase):
    def test_record_nunca_rompe_el_flujo_por_un_error(self):
        # ``changes`` con un valor no serializable a JSON (un set) hace
        # fallar el INSERT: record() debe tragarse la excepción y no
        # propagarla, dejando la fila sin crear.
        record(category="users", action="user.create", changes={"bad": {1, 2, 3}})
        self.assertEqual(AuditLog.objects.count(), 0)

    def test_record_crea_una_fila_con_los_datos_esperados(self):
        user = User.objects.create_user(
            username="audit_actor@example.com",
            email="audit_actor@example.com",
            password=STRONG_PASSWORD,
        )
        record(
            actor=user,
            category=AuditLog.Category.USERS,
            action=AuditLog.Action.USER_CREATE,
            target=user,
            changes={"email": {"from": None, "to": user.email}},
        )
        log = AuditLog.objects.get()
        self.assertEqual(log.actor_id, user.id)
        self.assertEqual(log.actor_email, user.email)
        self.assertEqual(log.target_type, "user")
        self.assertEqual(log.target_id, str(user.pk))


class AuditFixtures(APITestCase):
    """Un admin con ``audit.view``, un usuario sin permiso y dos registros."""

    def setUp(self):
        for name in ("admin", "subscriber"):
            Group.objects.get_or_create(name=name)

        self.admin_user = User.objects.create_user(
            username="audit_admin@example.com",
            email="audit_admin@example.com",
            password=STRONG_PASSWORD,
            is_staff=True,
        )
        self.admin_user.groups.add(Group.objects.get(name="admin"))

        self.plain_user = User.objects.create_user(
            username="audit_plain@example.com",
            email="audit_plain@example.com",
            password=STRONG_PASSWORD,
        )
        self.plain_user.groups.add(Group.objects.get(name="subscriber"))

        audit_perm, _ = RolePermission.objects.get_or_create(
            key="audit.view", defaults={"name": "Ver auditoría", "category": "audit"}
        )
        GroupRolePermission.objects.get_or_create(
            group=Group.objects.get(name="admin"), permission=audit_perm
        )

        AuditLog.objects.create(
            category=AuditLog.Category.AUTH,
            action=AuditLog.Action.AUTH_LOGIN_SUCCESS,
            actor=self.admin_user,
            actor_email=self.admin_user.email,
        )
        AuditLog.objects.create(
            category=AuditLog.Category.USERS,
            action=AuditLog.Action.USER_CREATE,
            actor=self.admin_user,
            actor_email=self.admin_user.email,
            target_repr="nuevo@example.com",
        )


class AuditEndpointsTests(AuditFixtures):
    def test_no_admin_recibe_403(self):
        self.client.force_authenticate(user=self.plain_user)
        resp = self.client.get("/api/v1/audit/logs/")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_recibe_200(self):
        self.client.force_authenticate(user=self.admin_user)
        resp = self.client.get("/api/v1/audit/logs/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["pagination"]["count"], 2)

    def test_filtro_category_y_search_combinan_con_and(self):
        self.client.force_authenticate(user=self.admin_user)
        resp = self.client.get(
            "/api/v1/audit/logs/", {"category": "users", "search": "nuevo"}
        )
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["pagination"]["count"], 1)
        self.assertEqual(resp.data["results"][0]["action"], "user.create")

        resp = self.client.get(
            "/api/v1/audit/logs/", {"category": "users", "search": "no-matchea-nada"}
        )
        self.assertEqual(resp.data["pagination"]["count"], 0)

    def test_actions_catalogo(self):
        self.client.force_authenticate(user=self.admin_user)
        resp = self.client.get("/api/v1/audit/actions/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("categories", resp.data)
        self.assertIn("actions", resp.data)
        keys = {a["key"] for a in resp.data["actions"]}
        self.assertIn("user.create", keys)


class AuditExportTests(AuditFixtures):
    """GET /api/v1/audit/logs/export/ — los mismos filtros que el listado."""

    URL = "/api/v1/audit/logs/export/"

    def _csv_rows(self, resp):
        text = resp.content.decode("utf-8")
        self.assertTrue(text.startswith("﻿"), "el CSV tiene que traer BOM para Excel")
        return list(csv.reader(io.StringIO(text.lstrip("﻿")), delimiter=";"))

    def test_sin_permiso_403(self):
        self.client.force_authenticate(user=self.plain_user)
        self.assertEqual(self.client.get(self.URL).status_code, status.HTTP_403_FORBIDDEN)

    def test_csv_con_los_filtros_del_listado(self):
        self.client.force_authenticate(user=self.admin_user)
        resp = self.client.get(self.URL, {"category": "users", "search": "nuevo"})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("auditoria-", resp["Content-Disposition"])
        self.assertEqual(resp["X-Audit-Count"], "1")
        self.assertEqual(resp["X-Audit-Truncated"], "0")

        rows = self._csv_rows(resp)
        self.assertEqual(rows[0][0], "Fecha")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], self.admin_user.email)
        self.assertEqual(rows[1][6], "nuevo@example.com")

    def test_cambios_legibles_y_formula_neutralizada(self):
        AuditLog.objects.create(
            category=AuditLog.Category.USERS,
            action=AuditLog.Action.USER_CREATE,
            actor_email="=HYPERLINK(\"http://malo\")",
            target_repr="x",
            changes={"status": {"from": "active", "to": "inactive"}, "email": {"from": None, "to": "a@b.com"}},
        )
        self.client.force_authenticate(user=self.admin_user)
        rows = self._csv_rows(self.client.get(self.URL, {"search": "HYPERLINK"}))
        self.assertEqual(len(rows), 2)
        self.assertTrue(rows[1][1].startswith("'="), "una fórmula no puede llegar viva a la planilla")
        self.assertIn("status: active → inactive", rows[1][7])
        self.assertIn("email: — → a@b.com", rows[1][7])

    def test_xlsx(self):
        self.client.force_authenticate(user=self.admin_user)
        resp = self.client.get(self.URL, {"file_type": "xlsx"})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        sheet = load_workbook(io.BytesIO(resp.content)).active
        self.assertEqual(sheet.title, "Auditoría")
        self.assertEqual(sheet.max_row, 3)

    def test_tipo_de_archivo_invalido_400(self):
        self.client.force_authenticate(user=self.admin_user)
        resp = self.client.get(self.URL, {"file_type": "pdf"})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_se_corta_en_el_tope_y_lo_avisa(self):
        self.client.force_authenticate(user=self.admin_user)
        with patch("apps.audit.views.AUDIT_EXPORT_MAX_ROWS", 1):
            resp = self.client.get(self.URL)
        self.assertEqual(resp["X-Audit-Count"], "1")
        self.assertEqual(resp["X-Audit-Truncated"], "1")

"""Lectura de CSV/Excel para la importación de pedidos (story 21).

Python puro: el CSV se lee con el módulo ``csv`` de la stdlib, el ``.xlsx``
con ``openpyxl`` (no requiere ninguna librería de sistema, a diferencia de
otras opciones tipo pandas+xlrd). Ambos formatos terminan en la MISMA
forma: ``(headers: list[str], rows: list[dict[str, str]])``, para que
``import_views``/``ingestion`` no tengan que saber de qué formato vino el
archivo.
"""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook


class ImportFileError(Exception):
    """Archivo ilegible/vacío: se traduce a un 400 con el mensaje tal cual
    en la vista (nunca un 500 por un archivo mal formado que mandó el
    usuario)."""


def _decode_csv_bytes(raw_bytes):
    # Los Excel exportados en Argentina suelen salir en Latin-1 (Windows-1252
    # en la práctica) con ";" como separador de columnas, no en UTF-8 con ",".
    # "utf-8-sig" primero para tolerar el BOM que agrega Excel al guardar
    # "CSV UTF-8" a mano.
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("No se pudo leer la codificación del archivo (probá UTF-8 o Latin-1).")


def _sniff_delimiter(sample):
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
    except csv.Error:
        # Sniffer no siempre acierta con una sola columna o muy pocas filas:
        # cae al separador más frecuente en la muestra, con ";" como default
        # (más común en planillas argentinas) antes que "," en un empate.
        return ";" if sample.count(";") >= sample.count(",") else ","


def parse_csv(raw_bytes):
    text = _decode_csv_bytes(raw_bytes)
    delimiter = _sniff_delimiter(text[:4096])
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [row for row in reader if any((cell or "").strip() for cell in row)]
    if not rows:
        raise ImportFileError("El archivo está vacío.")

    headers = [(h or "").strip() for h in rows[0]]
    data_rows = []
    for raw_row in rows[1:]:
        data_rows.append(
            {
                headers[i]: (raw_row[i].strip() if i < len(raw_row) else "")
                for i in range(len(headers))
                if headers[i]
            }
        )
    return headers, data_rows


def parse_xlsx(file_obj):
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise ImportFileError("El archivo está vacío.")

        headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
        data_rows = []
        for raw_row in rows_iter:
            if raw_row is None or all(cell is None for cell in raw_row):
                continue
            data_rows.append(
                {
                    headers[i]: (
                        "" if i >= len(raw_row) or raw_row[i] is None else str(raw_row[i]).strip()
                    )
                    for i in range(len(headers))
                    if headers[i]
                }
            )
        return headers, data_rows
    finally:
        workbook.close()


def parse_tabular_file(django_file, filename):
    """``django_file``: un ``UploadedFile``/``FieldFile`` ya abierto en modo
    binario, posicionado al inicio. No lo cierra (el llamador decide, ver
    ``import_views``, que reabre un ``FieldFile`` para mapear/confirmar)."""
    if (filename or "").lower().endswith(".xlsx"):
        return parse_xlsx(django_file)
    return parse_csv(django_file.read())

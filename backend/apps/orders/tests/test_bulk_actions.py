"""Tests de las acciones masivas sobre pedidos propios (``bulk_views``):
cambio de estado, exportación, planilla de retiro y carga de seguimientos
desde la planilla del transportista."""

import csv
import io
import json
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from apps.audit.models import AuditLog
from apps.integrations.models import StoreConnection

from ..bulk import detect_columns, normalize_order_key
from ..models import Address, Order
from .test_orders import auth_headers_for

User = get_user_model()


class BulkTestBase(APITestCase):
    def setUp(self):
        subscriber, _ = Group.objects.get_or_create(name="subscriber")
        self.user = User.objects.create_user(username="vendedor1", email="vendedor1@example.com", password="Clave123!")
        self.user.groups.add(subscriber)
        self.other = User.objects.create_user(username="vendedor2", email="vendedor2@example.com", password="Clave123!")
        self.other.groups.add(subscriber)

        self.store_a = StoreConnection.objects.create(
            owner=self.user,
            platform=StoreConnection.Platform.TIENDANUBE,
            external_store_id="111",
            name="Tienda A",
            sender_name="Remitente A",
            sender_address="Av. Corrientes 1234, CABA",
        )
        self.store_b = StoreConnection.objects.create(
            owner=self.user, platform=StoreConnection.Platform.TIENDANUBE, external_store_id="222", name="Tienda B"
        )

        self.address = Address.objects.create(
            user=self.user,
            recipient_name="María Pérez",
            street="Calle Falsa",
            number="123",
            city="Rosario",
            state="Santa Fe",
            postal_code="2000",
            origin=Address.Origin.SHIPMENT,
        )
        self.order_a = Order.objects.create(
            user=self.user,
            address=self.address,
            store_connection=self.store_a,
            external_id="a1",
            external_number="1001",
            package_count=2,
            total_weight_kg=Decimal("1.500"),
            items=[{"name": "Remera", "quantity": 2}],
        )
        self.order_b = Order.objects.create(
            user=self.user, address=self.address, store_connection=self.store_b, external_id="b1", external_number="2001"
        )
        self.manual = Order.objects.create(user=self.user, address=self.address)

        other_address = Address.objects.create(user=self.other, street="Otra", number="2", city="CABA")
        self.foreign = Order.objects.create(user=self.other, address=other_address, external_number="1001")

    def post(self, url, payload, user=None, **kwargs):
        return self.client.post(url, payload, format=kwargs.pop("format", "json"), **auth_headers_for(user or self.user), **kwargs)


class BulkStatusTests(BulkTestBase):
    URL = "/api/v1/orders/bulk-status/"

    def test_despacha_varios_y_audita_cada_uno(self):
        response = self.post(
            self.URL,
            {"order_ids": [self.order_a.pk, self.order_b.pk], "status": "dispatched", "carrier": "Correo"},
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["updated"], 2)
        self.assertEqual(response.data["failed"], 0)
        for order in (self.order_a, self.order_b):
            order.refresh_from_db()
            self.assertEqual(order.status, Order.Status.DISPATCHED)
            self.assertEqual(order.carrier, "Correo")
        self.assertEqual(AuditLog.objects.filter(action="order.ship", actor=self.user).count(), 2)
        self.assertFalse(AuditLog.objects.filter(action="order.status_change").exists())

    def test_un_pedido_que_no_puede_avanzar_no_corta_el_resto(self):
        self.order_a.status = Order.Status.DELIVERED
        self.order_a.save()

        response = self.post(self.URL, {"order_ids": [self.order_a.pk, self.manual.pk], "status": "in_transit"})

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["updated"], 1)
        failed = [item for item in response.data["results"] if not item["ok"]]
        self.assertEqual([item["order_id"] for item in failed], [self.order_a.pk])
        self.assertEqual(failed[0]["number"], "1001")
        self.manual.refresh_from_db()
        self.assertEqual(self.manual.status, Order.Status.IN_TRANSIT)

    def test_no_retrocede(self):
        self.order_a.status = Order.Status.IN_TRANSIT
        self.order_a.save()
        response = self.post(self.URL, {"order_ids": [self.order_a.pk], "status": "dispatched"})
        self.assertEqual(response.data["failed"], 1)
        self.order_a.refresh_from_db()
        self.assertEqual(self.order_a.status, Order.Status.IN_TRANSIT)

    def test_cancela_solo_los_cancelables(self):
        self.order_b.status = Order.Status.DISPATCHED
        self.order_b.save()

        response = self.post(self.URL, {"order_ids": [self.order_a.pk, self.order_b.pk], "status": "cancelled"})

        self.assertEqual(response.data["updated"], 1)
        self.order_a.refresh_from_db()
        self.order_b.refresh_from_db()
        self.assertEqual(self.order_a.status, Order.Status.CANCELLED)
        self.assertEqual(self.order_b.status, Order.Status.DISPATCHED)
        self.assertTrue(AuditLog.objects.filter(action="order.cancel", target_id=str(self.order_a.pk)).exists())

    def test_pedido_ajeno_se_informa_como_no_encontrado_y_no_se_toca(self):
        response = self.post(self.URL, {"order_ids": [self.foreign.pk], "status": "dispatched"})
        self.assertEqual(response.data["failed"], 1)
        self.foreign.refresh_from_db()
        self.assertEqual(self.foreign.status, Order.Status.CREATED)

    def test_lista_vacia_y_estado_invalido(self):
        self.assertEqual(self.post(self.URL, {"order_ids": [], "status": "dispatched"}).status_code, 400)
        self.assertEqual(self.post(self.URL, {"order_ids": [self.manual.pk], "status": "created"}).status_code, 400)

    def test_requiere_autenticacion(self):
        response = self.client.post(self.URL, {"order_ids": [self.manual.pk], "status": "dispatched"}, format="json")
        self.assertEqual(response.status_code, 401)


class OrderExportTests(BulkTestBase):
    URL = "/api/v1/orders/export/"

    def _get(self, query=""):
        return self.client.get(f"{self.URL}{query}", **auth_headers_for(self.user))

    def _csv_rows(self, response):
        text = response.content.decode("utf-8-sig")
        return list(csv.reader(io.StringIO(text), delimiter=";"))

    def test_csv_con_los_pedidos_propios(self):
        response = self._get()

        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response["Content-Disposition"])
        rows = self._csv_rows(response)
        self.assertEqual(rows[0][0], "Pedido")
        numbers = [row[0] for row in rows[1:]]
        self.assertEqual(numbers, ["1001", "2001", str(self.manual.pk)])
        first = dict(zip(rows[0], rows[1]))
        self.assertEqual(first["Destinatario"], "María Pérez")
        self.assertEqual(first["Tienda"], "Tienda A")
        self.assertEqual(first["Peso (kg)"], "1.5")
        self.assertEqual(first["Productos"], "2 x Remera")
        self.assertEqual(response["X-Order-Count"], "3")

    def test_respeta_los_filtros_del_listado(self):
        rows = self._csv_rows(self._get(f"?store={self.store_b.pk}"))
        self.assertEqual([row[0] for row in rows[1:]], ["2001"])
        rows = self._csv_rows(self._get(f"?ids={self.manual.pk}"))
        self.assertEqual(len(rows), 2)

    def test_excel(self):
        response = self._get("?file_type=xlsx")
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet["A1"].value, "Pedido")
        self.assertEqual(sheet.max_row, 4)

    def test_neutraliza_formulas_de_texto_de_terceros(self):
        self.address.recipient_name = "=HYPERLINK(\"http://x\")"
        self.address.save()
        rows = self._csv_rows(self._get())
        self.assertTrue(rows[1][5].startswith("'="))

        sheet = load_workbook(io.BytesIO(self._get("?file_type=xlsx").content)).active
        self.assertEqual(sheet["F2"].data_type, "s")

    def test_tipo_invalido(self):
        self.assertEqual(self._get("?file_type=pdf").status_code, 400)


class DispatchManifestTests(BulkTestBase):
    URL = "/api/v1/orders/manifest/"

    def test_devuelve_un_pdf(self):
        response = self.post(self.URL, {"order_ids": [self.order_a.pk, self.manual.pk], "carrier": "Correo <b>"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_muchos_pedidos_ocupan_varias_paginas(self):
        orders = [
            Order.objects.create(user=self.user, address=self.address, external_id=f"m{i}") for i in range(80)
        ]
        response = self.post(self.URL, {"order_ids": [order.pk for order in orders]})
        self.assertEqual(response.status_code, 200)
        self.assertGreater(response.content.count(b"/Type /Page\n"), 1)

    def test_solo_pedidos_ajenos_da_400(self):
        response = self.post(self.URL, {"order_ids": [self.foreign.pk]})
        self.assertEqual(response.status_code, 400)


def _csv_file(rows, name="seguimientos.csv"):
    buffer = io.StringIO()
    csv.writer(buffer, delimiter=";").writerows(rows)
    return SimpleUploadedFile(name, buffer.getvalue().encode("utf-8"), content_type="text/csv")


def _xlsx_file(rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile("guias.xlsx", buffer.getvalue())


class TrackingImportTests(BulkTestBase):
    PREVIEW = "/api/v1/orders/tracking-import/preview/"
    CONFIRM = "/api/v1/orders/tracking-import/confirm/"

    def _preview(self, upload, **extra):
        return self.post(self.PREVIEW, {"file": upload, **extra}, format="multipart")

    def test_detecta_columnas_y_cruza_por_numero_de_tienda(self):
        upload = _csv_file(
            [
                ["Nro de pedido", "Número de seguimiento", "Transportista"],
                ["#1001", "AR111", "Correo"],
                ["2001", "AR222", ""],
                [str(self.manual.pk), "AR333", "Moto"],
                ["9999", "AR444", ""],
            ]
        )

        response = self._preview(upload, carrier="Andreani")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(
            response.data["mapping"],
            {
                "order": "Nro de pedido",
                "tracking_number": "Número de seguimiento",
                "carrier": "Transportista",
                "tracking_url": None,
            },
        )
        rows = response.data["rows"]
        self.assertEqual([row["state"] for row in rows], ["ok", "ok", "ok", "not_found"])
        self.assertEqual(rows[0]["order"]["id"], self.order_a.pk)
        # Sin transportista en la fila, vale el de toda la planilla.
        self.assertEqual(rows[1]["carrier"], "Andreani")
        self.assertEqual(rows[2]["order"]["id"], self.manual.pk)
        self.assertEqual(rows[3]["row"], 5)
        self.assertEqual(response.data["counts"], {"ok": 3, "not_found": 1})

    def test_nunca_cruza_con_un_pedido_ajeno(self):
        # El otro vendedor también tiene un #1001: solo debe aparecer el propio.
        response = self._preview(_csv_file([["pedido", "seguimiento"], ["1001", "X1"]]))
        self.assertEqual(response.data["rows"][0]["order"]["id"], self.order_a.pk)

    def test_numero_repetido_entre_tiendas_es_ambiguo_y_la_tienda_desempata(self):
        self.order_b.external_number = "1001"
        self.order_b.save()
        upload = lambda: _csv_file([["pedido", "seguimiento"], ["1001", "X1"]])  # noqa: E731

        response = self._preview(upload())
        self.assertEqual(response.data["rows"][0]["state"], "ambiguous")

        response = self._preview(upload(), store=str(self.store_b.pk))
        self.assertEqual(response.data["rows"][0]["state"], "ok")
        self.assertEqual(response.data["rows"][0]["order"]["id"], self.order_b.pk)

    def test_estados_de_fila_con_problemas(self):
        self.order_b.status = Order.Status.CANCELLED
        self.order_b.save()
        upload = _csv_file(
            [
                ["pedido", "seguimiento", "url"],
                ["1001", "", ""],
                ["2001", "X2", ""],
                ["", "X3", ""],
                [str(self.manual.pk), "X4", "no es una url"],
                [str(self.manual.pk), "X5", ""],
            ]
        )
        rows = self._preview(upload).data["rows"]
        self.assertEqual(
            [row["state"] for row in rows],
            ["missing_tracking", "not_shippable", "missing_order", "invalid_url", "duplicate"],
        )

    def test_excel_con_numeros_como_celdas_numericas(self):
        upload = _xlsx_file([["Pedido", "Guía"], [1001, "G-1"], [2001.0, "G-2"]])
        response = self._preview(upload)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual([row["state"] for row in response.data["rows"]], ["ok", "ok"])

    def test_columnas_elegidas_a_mano(self):
        upload = _csv_file([["col1", "col2"], ["1001", "Z9"]])

        response = self._preview(upload)
        self.assertEqual(response.data["rows"], [])
        self.assertEqual(response.data["headers"], ["col1", "col2"])

        upload = _csv_file([["col1", "col2"], ["1001", "Z9"]])
        response = self._preview(upload, mapping=json.dumps({"order": "col1", "tracking_number": "col2"}))
        self.assertEqual(response.data["rows"][0]["state"], "ok")

        upload = _csv_file([["col1", "col2"], ["1001", "Z9"]])
        response = self._preview(upload, mapping=json.dumps({"order": "nope"}))
        self.assertEqual(response.status_code, 400)

    def test_archivo_invalido(self):
        self.assertEqual(self._preview(SimpleUploadedFile("x.pdf", b"%PDF")).status_code, 400)
        self.assertEqual(self.post(self.PREVIEW, {}, format="multipart").status_code, 400)

    def test_confirmar_carga_seguimiento_y_despacha(self):
        response = self.post(
            self.CONFIRM,
            {
                "status": "dispatched",
                "rows": [
                    {"order_id": self.order_a.pk, "tracking_number": "AR111", "carrier": "Correo",
                     "tracking_url": "https://seguimiento.example.com/AR111"},
                    {"order_id": self.foreign.pk, "tracking_number": "AR999"},
                ],
            },
        )

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["updated"], 1)
        self.order_a.refresh_from_db()
        self.assertEqual(self.order_a.status, Order.Status.DISPATCHED)
        self.assertEqual(self.order_a.tracking_number, "AR111")
        self.assertEqual(self.order_a.carrier, "Correo")
        self.assertEqual(self.order_a.tracking_url, "https://seguimiento.example.com/AR111")
        self.foreign.refresh_from_db()
        self.assertEqual(self.foreign.tracking_number, "")
        log = AuditLog.objects.get(action="order.ship")
        self.assertEqual(log.changes["tracking_number"], {"from": "", "to": "AR111"})

    def test_confirmar_no_hace_retroceder_un_pedido_mas_avanzado(self):
        self.order_a.status = Order.Status.IN_TRANSIT
        self.order_a.save()
        response = self.post(
            self.CONFIRM,
            {"status": "dispatched", "rows": [{"order_id": self.order_a.pk, "tracking_number": "N1"}]},
        )
        self.assertEqual(response.data["updated"], 1)
        self.order_a.refresh_from_db()
        self.assertEqual(self.order_a.status, Order.Status.IN_TRANSIT)
        self.assertEqual(self.order_a.tracking_number, "N1")

    def test_confirmar_rechaza_url_invalida_por_fila(self):
        response = self.post(
            self.CONFIRM,
            {"status": "dispatched", "rows": [{"order_id": self.order_a.pk, "tracking_number": "N1", "tracking_url": "nada"}]},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["failed"], 1)
        self.order_a.refresh_from_db()
        self.assertEqual(self.order_a.tracking_number, "")


class TrackingHelpersTests(APITestCase):
    def test_detect_columns_no_confunde_seguimiento_con_pedido(self):
        mapping = detect_columns(["Número de seguimiento", "Número de pedido", "Empresa", "Link"])
        self.assertEqual(mapping["tracking_number"], "Número de seguimiento")
        self.assertEqual(mapping["order"], "Número de pedido")
        self.assertEqual(mapping["carrier"], "Empresa")
        self.assertEqual(mapping["tracking_url"], "Link")

    def test_ref_no_cae_en_otra_palabra(self):
        self.assertIsNone(detect_columns(["Preferencia"])["order"])

    def test_normalize_order_key(self):
        self.assertEqual(normalize_order_key(" #1001 "), "1001")
        self.assertEqual(normalize_order_key("1001.0"), "1001")
        self.assertEqual(normalize_order_key(None), "")


class BulkMenuTests(BulkTestBase):
    def test_el_menu_ofrece_las_acciones_masivas(self):
        response = self.client.get("/api/v1/auth/users/me/dashboard/", **auth_headers_for(self.user))
        urls = {item["key"]: item["url"] for item in response.data["menu"]}
        # La planilla de retiro y cargar seguimientos se hacen desde Mis
        # pedidos, sin botón propio (decidido 2026-10-09).
        self.assertNotIn("dispatch_manifest", urls)
        self.assertNotIn("tracking_import", urls)
        self.assertEqual(urls["orders_bulk_status"], "estado_pedidos.html")
        self.assertEqual(urls["orders_export"], "exportar_pedidos.html")
